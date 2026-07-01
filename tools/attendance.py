import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl import load_workbook
from copy import copy
from datetime import date
import calendar
import re
import os

def process_attendance(old_path, template_path):
    # 1. 读取原始表，提取统计日期
    old_raw = pd.read_excel(old_path, header=None)
    title_row = old_raw.iloc[0, 0]
    date_match = re.search(r'统计日期：(\d{4})-(\d{2})-\d{2}', str(title_row))
    if date_match:
        year = int(date_match.group(1))
        month = int(date_match.group(2))
    else:
        year = 2025
        month = 11

    # ===== 核心修改：根据月份计算实际天数 =====
    _, days_in_month = calendar.monthrange(year, month)  # 实际天数，如2月为28/29

    # 生成日期和星期（只生成实际天数）
    week_days = []
    dates = []
    for day in range(1, days_in_month + 1):
        try:
            d = date(year, month, day)
        except:
            week_days.append("")
            dates.append("")
            continue
        weekday = d.weekday()
        week_map = ["一", "二", "三", "四", "五", "六", "日"]
        week_days.append(week_map[weekday])
        dates.append(day)

    # 读取原始表加班汇总列：AC=工作日加班 AD=休息日加班 AE=节假日加班
    header_row = old_raw.iloc[3]
    workday_overtime_col = None
    rest_overtime_col = None
    holiday_overtime_col = None
    for idx, col_name in enumerate(header_row):
        if pd.isna(col_name):
            continue
        col_str = str(col_name).strip()
        if "工作日加班" in col_str:
            workday_overtime_col = idx
        elif "休息日加班" in col_str:
            rest_overtime_col = idx
        elif "节假日加班" in col_str:
            holiday_overtime_col = idx

    # 处理人员数据
    data_rows_old = old_raw.iloc[4:]
    name_col = 0
    # ===== 核心修改：日期列只取实际天数 =====
    daily_cols_old = list(range(31, 31 + days_in_month))  # 例：30天为31~60，31天为31~61
    old_people = {}
    for idx, row in data_rows_old.iterrows():
        name = row[name_col]
        if pd.isna(name):
            continue
        name = str(name).strip()
        daily_status = []
        work_days = 0
        leave_days = 0

        # 读取加班数据（强制转数字）
        workday_overtime = float(row[workday_overtime_col]) if (workday_overtime_col is not None and pd.notna(row[workday_overtime_col])) else 0.0
        rest_overtime = float(row[rest_overtime_col]) if (rest_overtime_col is not None and pd.notna(row[rest_overtime_col])) else 0.0
        holiday_overtime = float(row[holiday_overtime_col]) if (holiday_overtime_col is not None and pd.notna(row[holiday_overtime_col])) else 0.0

        overtime_days = rest_overtime + holiday_overtime
        night_shift = workday_overtime * 2  # 夜班次=工作日加班×2

        # 每日考勤状态（只遍历实际存在的列）
        for i, col in enumerate(daily_cols_old):
            # 已无 break，因为列数正好等于天数
            status = row[col]
            if pd.isna(status):
                daily_status.append(None)
                continue
            status_str = str(status).strip()

            if "产假" in status_str:
                sym = "产"
                daily_status.append(sym)
                continue
            elif "0.5天" in status_str:
                if "休息" in status_str:
                    sym = "■"
                    daily_status.append(sym)
                elif "事假" in status_str or "病假" in status_str:
                    work_days += 0.5
                    leave_days += 0.5
                    sym = "●"
                    daily_status.append(sym)
                else:
                    sym = "√"
                    daily_status.append(sym)
                    work_days += 1
                continue
            elif "病假" in status_str:
                sym = "△"
                daily_status.append(sym)
                leave_days += 1
                continue
            elif "事假" in status_str:
                sym = "○"
                daily_status.append(sym)
                leave_days += 1
                continue
            elif "加班" in status_str:
                if "休息" in status_str:
                    if "13:00" in status_str:
                        sym = "■"
                        daily_status.append(sym)
                    else:
                        sym = "□"
                        daily_status.append(sym)
                else:
                    sym = "√"
                    work_days += 1
                    daily_status.append(sym)
                continue
            elif "休息" in status_str:
                daily_status.append(None)
                continue
            elif "住院假" in status_str:
                sym = "住"
                daily_status.append(sym)
                leave_days += 1
                continue
            elif "丧假" in status_str:
                sym = "丧"
                daily_status.append(sym)
                work_days += 1
                continue
            elif "产检" in status_str:
                sym = "产检"
                daily_status.append(sym)
                work_days += 1
                continue
            elif "婚假" in status_str:
                sym = "婚"
                daily_status.append(sym)
                work_days += 1
                continue
            elif "正常" in status_str:
                sym = "√"
                daily_status.append(sym)
                work_days += 1
                continue
            else:
                sym = "√"
                daily_status.append(sym)
                work_days += 1
                continue

        old_people[name] = {
            "daily": daily_status,
            "work_days": work_days,
            "overtime_days": overtime_days,
            "leave_days": leave_days,
            "night_shift": night_shift
        }

    # 加载底板
    wb = load_workbook(template_path)
    ws = wb.active
    uni_font = Font(name="宋体", size=10)
    uni_align = Alignment(horizontal="center", vertical="center")

    # 提取部门名
    dept_name = ""
    dept_cell = ws.cell(row=2, column=1).value
    if dept_cell:
        dept_name = str(dept_cell).strip()
        dept_name = re.sub(r'\d{4}年|\d{1,2}月', '', dept_name)
        dept_name = dept_name.replace("考勤表", "").replace("考勤", "").replace("部门：", "").strip()

    # 填充AI2日期
    target_row = 2
    target_col = 35
    target_range = None
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= target_row <= merged_range.max_row and merged_range.min_col <= target_col <= merged_range.max_col:
            target_range = merged_range
            break
    date_text = f"{year}年{month}月"
    if target_range:
        ws.cell(row=target_range.min_row, column=target_range.min_col, value=date_text)
    else:
        ws.cell(row=target_row, column=target_col, value=date_text)

    # ===== 更新表头：只更新实际天数对应的列（其余保留原样） =====
    date_row_idx = 4
    for i, d in enumerate(dates):
        col = 3 + i
        cell = ws.cell(row=date_row_idx, column=col, value=d)
        cell.font = uni_font
        cell.alignment = uni_align
    # 对于多余的天数（模板可能预填了31天），可以清空或保持，但不清除也可
    # 若想清空多余列，可循环 from days_in_month to 30（0-based）
    for i in range(days_in_month, 31):  # 31是固定上限
        col = 3 + i
        cell = ws.cell(row=date_row_idx, column=col, value=None)
        cell.font = uni_font
        cell.alignment = uni_align

    week_row_idx = 5
    for i, w in enumerate(week_days):
        col = 3 + i
        cell = ws.cell(row=week_row_idx, column=col, value=w)
        cell.font = uni_font
        cell.alignment = uni_align
    for i in range(days_in_month, 31):
        col = 3 + i
        cell = ws.cell(row=week_row_idx, column=col, value=None)
        cell.font = uni_font
        cell.alignment = uni_align

    # 检测AK3单元格是否为“夜班次”
    night_shift_col = None
    ak_col = 37  # AK列固定是第37列
    ak3_val = ws.cell(row=3, column=ak_col).value  # 检测AK3单元格
    if pd.notna(ak3_val) and "夜班次" in str(ak3_val):
        night_shift_col = ak_col  # 是夜班次，则填充到AK列

    # 读取模板人员名单
    template_names = []
    start_row = 6
    for row_idx in range(start_row, ws.max_row + 1):
        name = ws.cell(row=row_idx, column=2).value
        if pd.isna(name):
            continue
        template_names.append((row_idx, str(name).strip()))

    # 填充数据
    for row_idx, name in template_names:
        if name not in old_people:
            continue
        person = old_people[name]

        # 每日符号（只填充实际天数）
        for i, sym in enumerate(person["daily"]):  # 此时列表长度就是实际天数
            col = 3 + i
            cell = ws.cell(row=row_idx, column=col, value=sym)
            cell.font = uni_font
            cell.alignment = uni_align
            cell.border = copy(ws.cell(start_row, col).border)
        # 多余天数的格子可清空（保持空白）
        for i in range(len(person["daily"]), 31):
            col = 3 + i
            cell = ws.cell(row=row_idx, column=col, value=None)
            cell.font = uni_font
            cell.alignment = uni_align
            cell.border = copy(ws.cell(start_row, col).border)

        # AH列（出勤）
        val = person["work_days"]
        v = "" if (val == 0 or val == 0.0) else (int(val) if val.is_integer() else val)
        cell = ws.cell(row_idx, 34, value=v)
        cell.font = uni_font
        cell.alignment = uni_align
        cell.border = copy(ws.cell(start_row, 34).border)
        cell.number_format = "General"

        # AI列（加班）
        val = person["overtime_days"]
        v = "" if (val == 0 or val == 0.0) else (int(val) if val.is_integer() else val)
        cell = ws.cell(row_idx, 35, value=v)
        cell.font = uni_font
        cell.alignment = uni_align
        cell.border = copy(ws.cell(start_row, 35).border)
        cell.number_format = "General"

        # AJ列（请假）
        val = person["leave_days"]
        v = "" if (val == 0 or val == 0.0) else (int(val) if val.is_integer() else val)
        cell = ws.cell(row_idx, 36, value=v)
        cell.font = uni_font
        cell.alignment = uni_align
        cell.border = copy(ws.cell(start_row, 36).border)
        cell.number_format = "General"

        # AK列（夜班次）
        if night_shift_col is not None:
            val = person["night_shift"]
            v = "" if (val == 0 or val == 0.0) else (int(val) if val.is_integer() else val)
            cell = ws.cell(row=row_idx, column=night_shift_col, value=v)
            cell.font = uni_font
            cell.alignment = uni_align
            cell.border = copy(ws.cell(start_row, night_shift_col).border)
            cell.number_format = "General"

    # 列宽
    ws.column_dimensions['AH'].width = 8
    ws.column_dimensions['AI'].width = 8
    ws.column_dimensions['AJ'].width = 8
    if night_shift_col is not None:
        ws.column_dimensions['AK'].width = 8

    # 保存
    output_name = f"{year}年{month}月{dept_name}考勤.xlsx"
    output_path = os.path.join(os.path.dirname(old_path), output_name)
    wb.save(output_path)
    return output_path, output_name